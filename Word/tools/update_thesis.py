from pathlib import Path
import hashlib, json, re
from copy import deepcopy
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[2]
WORK = ROOT / 'Word/materials/thesis_update'
SOURCE = ROOT / 'Word/本科毕业论文_基于TensorFlow的智能电网负荷预测系统设计与实现_初始化.docx'
OUTPUT = ROOT / 'Word/本科毕业论文_基于TensorFlow的智能电网负荷预测系统设计与实现_项目更新排版版.docx'
source_sha = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
doc = Document(SOURCE)
original = list(doc.paragraphs)
tables = list(doc.tables)
changes = []

def replace(i, text):
    changes.append({'source_paragraph': i, 'old': original[i].text, 'new': text})
    original[i].text = text

updates = {
4: '基于TensorFlow的智能电网\n负荷预测系统设计与实现',
8: '面向ISO New England（ISO-NE）控制区的短期运行分析，本文设计并实现基于TensorFlow的智能电网负荷预测系统，并扩展电价分位预测与分布式光伏预测。三个独立模型分别以168小时负荷历史、168小时电价历史和96小时光伏历史为输入，结合未来天气与日历特征，输出未来24小时结果。系统采用React、TypeScript、FastAPI与MySQL，在Windows本地运行。在线流程统一使用美国东部时区，区分小时开始与小时结束标签；通过启动预加载、共享请求和有效缓存恢复改善首次数据展示，并分别报告三个任务的数据状态。系统新增在线负荷提前量分组、实际输入归档及24步离线重放，使生成时间、数据来源和评价对象可以追溯。固定测试集上，负荷MAE为305.40 MW、RMSE为476.14 MW、MAPE为1.922%、R²为0.97739；电价P50 MAE为13.266 USD/MWh，P10—P90覆盖率为66.93%；光伏全时段MAE为236.1834 MW，白天WAPE为17.0532%。光伏指标以ISO-NE估算出力为参考。负荷和电价固定测试使用事后气象，不能直接代表严格在线预报精度；估计输入不作为真实标签。现有模型权重保持不变，新增电价区间校准候选未通过验证。阶段验收确认本地启动、重启缓存、真实接口与浏览器交互可用，长期稳定性、极端场景精度和夏令时完整日处理仍有待验证。',
12: 'This thesis develops a TensorFlow-based load forecasting system for the ISO New England (ISO-NE) control area, with additional electricity-price quantiles and distributed photovoltaic (PV) forecasts. Three independent models combine historical sequences and future weather and calendar features to forecast 24 hours ahead. Load and price use 168-hour histories; PV uses a 96-hour history. React, TypeScript, FastAPI, and MySQL support local operation on Windows. The online pipeline uses the America/New_York time zone and distinguishes hour-start PV labels from hour-ending load and price labels. Startup preloading, shared requests, and validated cache restoration improve initial data availability. Each task reports its own availability and input quality. New features include lead-time groups for online load errors, actual-input archives, and offline replay of all 24 forecast steps. On the fixed test sets, load MAE is 305.40 MW, RMSE is 476.14 MW, MAPE is 1.922%, and R² is 0.97739. Price P50 MAE is 13.266 USD/MWh, with P10–P90 coverage of 66.93%. PV all-period MAE is 236.1834 MW and daytime WAPE is 17.0532%, relative to ISO-NE estimated behind-the-meter PV. Fixed load and price tests use retrospective weather observations and do not establish strict online forecast accuracy. Estimated inputs are excluded from reference labels. Production weights remain unchanged, and two price-interval calibration candidates failed validation. Staged acceptance records confirm local startup, restart recovery, live API access, and browser interaction. Long-term reliability, extreme-event accuracy, and complete daylight-saving-day handling require further evaluation.',
113: '技术路线按数据、模型、服务与评价四个层次展开。数据层核对ISO-NE来源、时间区间和缺值；模型层按目标时间划分训练、验证和测试，训练分区拟合缩放参数，避免预测目标进入未来输入；服务层组织并行资料准备、独立任务状态和有效缓存；评价层区分固定离线测试、事后天气回测、在线提前量统计与原输入重放。上述措施控制目标和分区泄漏，但既有事后天气与切分前插值仍限制严格在线可用性结论。',
109: '（4）完成预测服务与数据管理。FastAPI加载冻结资产并协调共享快照，MySQL按来源保存预测与完整标签；内存及本地JSON承担预测缓存，压缩归档保留实际输入。估计补值不作为实测标签，含估计输入的预测与完整标签配对时单独分组评价。Redis和Prometheus/Grafana作为可选配置保留。',
124: '多步预测需要区分点预测和区间预测。负荷模型与PV v2输出单一数值序列，电价模型输出P10、P50和P90。分位有序不等于区间已经校准：校准参数应在验证数据上拟合和筛选，独立测试集仅用于最终评价，不应依据测试误差反向选择修正方案。',
150: '系统通过Windows PowerShell脚本在本机启动前后端，Docker Compose与监控配置作为可选资料保留。用户确认的开发环境为Windows 11、WSL2 Ubuntu 22.04、Python 3.10和RTX 3050 Laptop 4GB，TensorFlow能够识别GPU。本地用于预处理、小规模形状及梯度检查、保存加载和推理；正式长时间或多模型训练优先使用Google Colab T4。该环境约定不等于现有冻结权重的历史训练硬件记录；现存训练日志未完整保存GPU型号、CUDA和cuDNN版本。',
148: '数据持久化使用MySQL、SQLAlchemy及aiomysql/PyMySQL，分别保存预测、负荷、天气、性能、日志和系统指标。当前共享缓存为内存与本地JSON；Redis与监控工具仅保留可选配置。数据库在线累计结果与冻结实验文件分别管理，防止混用评价口径。',
158: '离线PV v2数据脚本与在线特征提供器读取ISO-NE五分钟分区估算负荷接口，使用八个分区的interval_begin_date、load_zone_id、estimated_btm_pv_mw和estimated_load_mw字段[8,9]。这些数据的官方名称包含estimated，光伏目标属于估算参考出力。离线构建与在线评价的采样完整性规则不同，分别说明如下。',
168: 'PV v2离线天气来自Open-Meteo Previous Runs API的day-1字段。官方定义中，previous_day1对应有效时刻提前24小时的预报值[10]。脚本按八区坐标读取短波辐射、云量、温度和露点，并形成America/New_York小时开始索引。该接口按有效时间对齐固定提前量，不代表一个任意在线预测原点使用的24个天气值都来自同一已保存发行批次。',
169: '天气与功率按ts_start索引合并。构建脚本先对天气执行limit=2线性插值，再前向及反向填充；剩余缺失会终止构建。这些操作发生在模型时间切分之前，故不能据此断言每个填充值仅依赖预测原点当时可获得的信息；严格在线重演还需要逐条核验输入可用时刻。',
178: '离线PV v2构建按每区至少10个五分钟点筛选，并对天气执行有限插值及前后填充，对pv_mw_ISONE执行limit=1插值，残余缺失会终止构建。源码未将目标插值严格限定为已确认的春季跳时缺口，本文不作此推断。在线流程另设每区至少8点可作有标记输入、每区12点才可作完整小时评价标签的规则；补值、部分采样与未结束小时不得充当真实标签。',
191: 'PV v2天气使用day-1历史预报字段，负荷与电价固定测试使用回顾性SMD温度和露点，指定日期回测可能使用事后天气。三者具有不同信息条件。PV数据的切分前填充与按有效时间拼接，也尚不足以证明任意预测原点的全部未来特征在当时已经发布。',
186: '表3-2只列出训练脚本实际使用的输入维度，不把ca_features.parquet中未被模型读取的辅助列计入网络输入。',
195: '目标与分区泄漏控制包括：完整24步目标属于同一时间分区才保留窗口；缩放器仅在训练分区拟合；未来输入排除未来负荷、电价和光伏目标；目标滞后与滚动统计使用过去观测；训练shuffle只调整训练窗口顺序，不改变时间边界。仍需保留两项限制：负荷、电价的未来天气是事后观测；PV天气前后填充与目标插值先于切分。因此，当前实验应表述为采取了目标与分区隔离措施，而非完成全链路发行时刻无泄漏验证。',
212: 'BiGRU仅对预测原点之前的历史窗口进行双向编码，不读取未来目标；未来分支接收天气和日历等协变量。网络结构与目标隔离一致，但协变量能否在预测发行时刻获得，还取决于数据提供器与预报档案，不能单凭编码器结构证明。',
255: '生产调用链运行三个独立TensorFlow模型。近期工程更新集中于在线数据准备、时间对齐、缓存恢复、评价留档与前端展示，未重新训练或替换冻结权重。输入归档保存实际特征及模型配置、资产路径、大小和修改时间形成的签名；该签名用于识别资产变化，不等同于每次推理重新计算权重文件的内容哈希。离线诊断另外核对来源文件SHA-256。',
261: '本章依据实际网络定义和生产资产说明三个独立模型的输入、目标变换、损失与形状校验。现有权重以.weights.h5保存并由相同模型定义加载；后续正式训练拟将新模型、history和metrics另存独立实验目录，优先使用Colab T4，不覆盖既有资产。旧训练硬件记录不完整与输入发行时刻可用性仍是复现边界。',
270: '系统采用本地前后端分离架构。React负责页面与图表，FastAPI负责认证、输入编排、模型调用与数据库访问，三个TensorFlow模型独立推理，MySQL保留预测快照和按来源隔离的完整标签。当前预测缓存实装为内存快照与本地JSON，实际输入另存压缩归档；Redis仅有配置和健康探测，不是该缓存链的必经组件。Docker、Prometheus和Grafana作为可选配置保留，不作为本轮本地运行的必需依赖。',
271: '如图5-1所示，浏览器经本地代理访问FastAPI；后端并行获取ISO-NE历史与Open-Meteo天气，再按任务构造特征并共享一次在线计算。结果分别写入有效缓存、输入归档与预测记录，完整小时标签独立入库供评价。图中的缓存和归档承担不同用途：前者减少重启空白，后者保留实际预测输入用于追溯。',
275: 'FastAPI生命周期初始化数据库、服务容器和模型后，主动触发在线预加载。气象与电网历史资料并行获取，多页面共享同一进行中的计算。无有效预测时请求最多等待8秒，仍未完成则返回首次准备状态，后台继续执行；取消页面请求不取消共享任务。模型就绪、服务可用和预测资料就绪分别判断，避免把模型已加载显示为所有业务数据正常。',
276: '重启缓存保存于backend/cache/live_forecast.json，采用原子写入。恢复前校验格式版本、模型配置及资产签名、原生成时间、来源、连续24个目标小时、有限值和分位顺序。仅恢复原生成时间6小时以内且仍覆盖有效目标的结果；过期、损坏或模型变化则重新准备。返回缓存始终保留原生成时间与cached标记，不重复入库为新预测。手动刷新与并发请求复用共享计算，显式用户输入预测接口仍独立处理。',
277: '电价与光伏页面从共享快照分别读取分位和逐小时出力，独立显示输入质量及异常原因。PV独立页采用hour-start，在合并曲线中加1小时与负荷hour-ending对齐。指定日期电价回测复用生产模型，获取完整历史特征；可选范围末日自动推进至美东昨日。缺官方小时RT-LMP时仍可返回预测，实测与误差留空并说明配对数；事后天气回测不冒充保存的在线预测。',
278: '历史分析原有总误差按每目标小时最新快照计算，新增在线负荷提前量0<Δ≤6、6<Δ≤12、12<Δ≤24小时三组，页面显示为1—6、7—12、13—24小时。Δ由原生成时间至目标小时结束的UTC实际时长确定；每组每目标小时选最早快照，并区分完整输入、估计补齐和输入质量未知。组间可共享目标小时，样本数不能相加。评价只匹配同来源完整标签，旧混合actual_load_mw不直接作为新口径实测。',
286: '表5-3整理前端调用与运行所需的主要API。准确率接口同时提供总体和提前量分组，准备状态和部分可用状态属于业务响应。实际输入重放作为本地离线脚本执行，不在表中虚构额外HTTP端点。为便于阅读，同一路由组的次要接口在单元格内分行列示。',
289: 'React Router组织综合总览、负荷、电价、光伏、气象、运行态势、历史分析与系统状态八个受保护页面，并按页面懒加载。ApiContext合并公共请求，api.ts统一错误转换。准备中每5秒重试，失败、缓存沿用或分项不可用时每35秒重试，正常状态每5分钟更新；页面隐藏时暂停，恢复可见后重新判断状态。日期查询只允许当前选择对应的请求更新结果、错误和加载状态，防止旧日期响应覆盖新页面。',
290: '前端提供正式工作台与深色监控两种主题。正式模式采用藏蓝导航、灰蓝主体和浅冷面板；监控模式使用深蓝背景及分层面板。曲线颜色与线型共同区分实测、在线预测、历史回测和光伏，坐标、网格与提示框随主题调整。主题偏好在挂载前从本地存储恢复，存储受限时仍可在会话中切换。总览合并重复状态，页面分别呈现生成时间、采集时间与目标范围；估计补值不作为实测标签，含估计输入的预测与完整标签配对时单独分组评价。',
291: '图5-2和图5-3为2026年10月3日本地运行的负荷页面实拍，分别展示正式工作台与深色监控模式。图中动态预测值仅说明界面和状态表达，不作为固定测试集指标。截图来自同一前端版本，主题切换只影响呈现，不改变模型、认证或数据口径。',
292: '图5-2 正式工作台模式的负荷预测页面（2026-10-03）',
296: '参数错误、令牌失效和权限不足保留相应HTTP错误；预测与统计业务接口继续要求Bearer JWT。favicon.ico和Chrome开发工具的精确探测路径仅对GET/HEAD返回204，减少无关401日志，不扩展业务白名单。外部资料或单任务失败时，在响应中报告原因，并仅沿用有效旧结果；缺少有效数据时明确不可用，不吞掉异常或生成示例曲线。',
300: '系统实现JWT保护、参数校验、日志和独立服务状态。完整标签按来源单独入库，估计补值与未结束小时不作为实测标签；含估计输入的预测单独标注并按质量分组评价。2026年10月的阶段验收已覆盖真实外部接口、MySQL、本地浏览器及故障恢复；其范围与单元回归分开记录，不能据单次验收推断长期在线稳定性。',
303: '本章保留既有训练日志与冻结模型指标，并补充截至2026年10月3日的专项诊断和系统验收。后端依赖固定TensorFlow 2.16.1和Keras 3.3.3，数据处理使用pandas、NumPy及StandardScaler。前端依赖声明包括React 18.2、TypeScript 5.2、Vite 4.5和Recharts 2.8；依赖声明不等于本章逐项重测的精确安装版本。',
304: '运行与训练环境分开说明。正式交付以Windows本地启动为准；用户提供的WSL2 Ubuntu 22.04、Python 3.10与RTX 3050 Laptop 4GB用于预处理、轻量验证及推理，正式长训练优先Colab T4。2026年10月3日已有诊断实际使用Windows Python 3.11、TensorFlow 2.16.1 CPU，以batch=32处理完整验证与测试窗口，离线评估流程耗时22.922秒，包含资产加载、分批推理、诊断及输出写入，不含TensorFlow导入启动。旧训练仅记录GPU:0，不能据此认定在T4或RTX 3050上完成。',
305: '阶段验证分批进行。10月1日启动专项记录64项后端、13项前端用例通过；标签与恢复专项记录68项后端、16项前端通过。10月3日输入留档与评估专项记录82项后端及离线用例、5项前端用例通过，双主题另有13项主题与导航用例通过。各批范围和版本存在重叠，不能相加为全量测试数；类型检查和生产构建通过。以上为既有记录复核，本次论文编辑未重新训练或重跑项目测试。',
306: '表6-1区分本地开发配置、实际诊断环境和后续正式训练约定。未记录的旧训练硬件参数保留为未知，避免将新环境说明倒推为历史实验条件。',
315: '负荷测试包含722个窗口、17328个展开目标点，对应745个不同目标小时；滑动窗口目标重叠，17328不代表独立小时数。MAE、RMSE、MAPE、Bias与R²按24步展开计算，MAPE分母为max(|actual|,1)，Bias定义为prediction−actual。原Peak MAE按测试真实负荷的75%分位筛选，区别于后续按训练分位设定的阈值。未来天气为事后SMD观测，以下指标不能直接外推为严格在线预报精度。',
318: '原项目实验记录在相同722个目标窗口上重算DA Demand、naive-24与naive-168基线。表6-4与图6-2保留该记录：DA Demand的MAE为759.34 MW、RMSE为956.63 MW、MAPE为5.227%，对应负荷模型相对降低59.78%、50.23%和63.23%。这些数值来自既有论文实验汇总，不是本次新增基线实验；本次核对未取得独立固化的基线结果JSON，完整复现仍需补齐对应脚本与输出。',
319: '2026年10月3日诊断已保存逐窗口预测、参考值、预测原点和24步目标数组，并核对9个唯一来源文件的SHA-256与大小。新增高温和节假日分层结果见6.8节；这些诊断复用原模型，不构成重新训练，也不应把数组重放一致性解释为预测误差下降。',
334: 'PV v2验证集1080个窗口展开为25920个目标点，对应1103个不同目标小时；测试集600个窗口展开为14400点，对应623个不同小时。白天定义为sun_up>0.5或ISO-NE估算参考出力>50 MW，WAPE为绝对误差之和除以参考出力之和，条件MAPE仅统计参考值>500 MW。表6-6、表6-7的能量和峰值指标按重叠24小时窗口计算，不等同于完整日指标；全部误差属于估算参考误差。',
347: '本系统区分四类评价：固定离线测试使用冻结数据和时间切分；指定日期回测以事后资料重新推理；在线统计配对当时保存的预测与后来发布的完整标签；原输入重放核验相同输入与资产是否复现相同输出。四类结果不能互换，历史回测尤其不能称为当时提前24小时的在线精度。',
348: '在线负荷只配对zonal_he_v2预测与region=NewEngland_zonal_HE、data_source=iso_ne_zonal_he_v2的同一目标小时完整标签，不把数据库累计行数作为独立样本数。新增提前量统计用原生成时间计算，分组保留输入质量；缺少显式输入质量标记的旧预测归为未知，不事后补造质量。2026年10月3日一次30日统计请求耗时494.24毫秒，三个组分别有17、16和28个配对目标，仅说明当次可计算范围，不是同一完整样本集的精度对照。',
349: '电价回测范围末日随美东日期自动推进；跨月目标按hour-ending对齐。2026年10月1日验收中，2026年9月30日按需回测返回24个官方小时RT-LMP，目标为9月30日01:00至10月1日00:00，耗时16.564秒。缺标签保留null和排除原因，不计算无配对误差。当前特征或目标窗口跨夏令时转换时明确拒绝回测，不能宣称完整支持23/25小时日。',
351: '系统测试包括限定范围回归与真实运行两层。回归覆盖五分钟去重、缺区、非法值、整点、缓存到期、模型签名变化、任务独立降级、认证和日期请求竞态；真实运行记录涵盖Windows PowerShell 5.1启动及重启、MySQL、ISO-NE/Open-Meteo获取、登录后浏览器页面和失败后恢复。隔离测试中的上游失败不等于长时间真实断网试验，二者分别记录。',
352: '表6-8按日期与范围列出已有通过结果，不累加专项测试数。一次原输入归档为146401字节、耗时159.53毫秒；重放负荷24×1、电价24×3和光伏24步结果，最大绝对差均为0。这仅验证该次输入及原资产可复现，不表示预测与未来真实值误差为0。业务接口无令牌返回401，合法令牌返回200，无关图标探测返回204。',
353: '前端实际验收覆盖1440×900桌面与390×844移动视口，包括两种主题、主要页面真实曲线、菜单键盘操作和日期切换。图5-2、图5-3使用该版本真实截图。启动专项中首次请求获得24点的单次观测由26.508秒变为13.550秒，二者外网条件不完全一致且预加载已提前执行，不能作为完整启动耗时的严格同条件对照；重启有效缓存一次返回0.229秒。一次浏览器请求中断后35.286秒恢复，说明重试机制生效，不保证所有外部故障的恢复时间。',
356: '负荷固定测试优于原记录中的同索引基线，但总体MAE不能代表极端场景。新增冻结诊断中，温度≥86 °F（30 °C）的1182个重叠目标点MAE为733.65 MW，节假日576点为538.15 MW，非节假日16752点为297.39 MW，见表6-9。低温≤32 °F样本为0，不能推断寒潮精度。该分层使用事后天气，尚未证明增加温度修正即可改善在线结果。',
357: '电价分位有序并未保证80%覆盖。新增校准将验证集分为前段1057窗口拟合与后段361窗口筛选，隔离23个跨段窗口，避免共享目标小时。全局扩宽和按提前量扩宽均未满足后段预设的覆盖接近80%且区间评分下降要求，见表6-10，未启用生产。独立测试仍报告原区间66.93%覆盖，P50误差13.266 USD/MWh；不能依据候选在测试集上的更高覆盖率反向选型。',
358: 'PV v2白天WAPE为17.0532%，相对同索引naive-24降低57.42%，但总量与峰形仍需改善。新增25个完整当地日（600个不同目标小时）诊断得到日能量相对误差15.5244%、峰值绝对误差633.8880 MW、峰时绝对偏移0.9200小时，见表6-11。该口径与原600个重叠窗口能量误差14.8323%不同。现有云量分组脚本同时纳入八区列及cloud_mean、cloud_std，尚未形成一致的八区均值定义，本文暂不引用其分组数值。',
359: '工程验收确认共享输入、缓存恢复、任务独立状态和失败恢复可用，冻结模型未变。首次展示与归档耗时均来自有限次观测，未完成持续压力测试，页面也未测量帧率。短缺口估计输入的精度尚需独立统计；长时间外部断供、完整夏令时日和冬季、负电价场景仍缺充分证据，故系统定位为本地研究原型。',
362: '本章保留固定模型结果，新增分层诊断、验证集校准试验及分阶段系统验收。原输入重放证明结果可追溯，启动与页面恢复记录证明相应流程可用，均不自动提高模型预测精度。在线提前量误差、事后天气回测和固定测试继续分开解释；长期稳定性、估计输入精度及完整夏令时日处理仍待验证。',
365: '本文完成ISO-NE数据整理、独立TensorFlow模型、Windows本地服务与前端展示。负荷与电价使用168小时历史预测24小时，PV v2使用96小时历史预测24小时，窗口按完整目标时间切分，缩放器只在训练分区拟合。业务时区统一为America/New_York并明确小时区间；对歧义时间采取排除或拒绝策略，尚未完成23/25小时在线日支持。既有负荷、电价事后天气和PV切分前填充限制了严格在线复现结论。',
369: 'PV v2全时段MAE为236.1834 MW、白天WAPE为17.0532%，原重叠窗口能量误差为14.8323%，相对naive-24白天WAPE降低57.42%。工程上完成预加载、原时间缓存恢复、请求合并、分项状态、输入归档、24步重放与双主题页面，阶段验收包含真实接口和浏览器。上述变化改善可用性与结果解释，但未替换模型，不构成离线精度提升。',
371: '现有负荷、电价固定测试主要覆盖2026年7月，PV测试覆盖8月中旬至9月上旬。新增高温和节假日诊断揭示部分误差集中，但低温样本为0，电价测试负价点为0；暴雪、飓风、冬季高峰及长期结构变化尚缺冻结评价。滑动目标点重叠也限制样本独立性，不宜直接按展开点数推断统计显著性。',
372: '负荷总体Bias为+91.49 MW，原Peak MAE为567.18 MW，高温组MAE达733.65 MW。当前已保存逐步数组，但需要按相同目标样本和输入质量进行更长期的提前量分析；新在线分组各自目标集合不同，不能直接将组间差异解释为预测步长的因果影响。估计补齐输入下的精度尚未充分积累。',
373: '电价原区间覆盖率66.93%，两个验证集扩宽候选未通过后段筛选，生产仍使用原分位。缺少统一测试目标的传统电价基线，且负价样本为0；尖峰和符号变化的稳健性不能由现有指标充分说明。后续应先扩展验证时段与公平基线，再判断是否采用校准或重新训练。',
374: 'PV目标为ISO-NE估算BTM出力，而非区域电表总量。白天WAPE和完整日能量、峰时指标表明辐照变化与峰形值得继续研究，但天气填充和云量分组定义仍需审计。已留档的实际在线输入可以重放，普通天气端点未暴露model_run，采集时刻不能冒充预报发行时刻；旧预测也无法事后补造原输入。',
375: '系统本地运行依赖MySQL以及外部ISO-NE、Open-Meteo资料，预测缓存采用内存和本地JSON。已有真实联网、重启与浏览器验收不替代长期连续运行或大并发验证。数据库无偏移DATETIME不能唯一表达秋季重复小时，当前采取排除或拒绝策略。短时补值不作为真实标签，含估计输入的预测按质量分组评价；长时断供会使相应任务不可用。',
377: '后续优先积累同口径在线预测与完整标签，结合现有输入归档，按原生成时间与真实提前量评价完整和估计输入。需要取得或核验天气发行批次，而不把采集时间当作可用时间；在新的滚动验证集上复核缺值处理，避免切分前填充造成信息条件不清。',
378: '模型改进以冻结诊断为起点。负荷重点检查高温、节假日与峰值；电价补充同信息条件基线及负价、尖峰样本，校准仅在验证集拟合筛选；PV先统一八区云量分组并核验峰时和完整日电量。任何修正均在独立测试上评价误差、区间宽度和评分，不因单一指标改善直接启用。',
379: '若需新增模型训练，本地Windows/WSL与RTX 3050 4GB负责预处理、小规模smoke test、形状、梯度、保存加载和推理验证；正式、长时间及多模型训练优先Colab T4。新实验另存.keras模型、history与metrics，并保留输入列、切分、随机种子及来源签名；未经验证不覆盖既有权重或训练数据。',
380: '工程后续重点是将无偏移小时键升级为可区分UTC时刻与美东偏移的标识，验证春季23小时和秋季25小时日；补充外部接口长期故障、连续运行与资源监测。当前已完成的真实联网与双主题浏览器验收继续作为阶段证据，新增能力应以相应代码和测试记录确认。',
}
for i, text in updates.items():
    replace(i, text)
original[218].text += '其中pₜ表示原始RT_LMP目标，zₜ为变换后的数值。'
replace(198, '本章明确ISO-NE与Open-Meteo数据来源、美东小时区间和采样门槛。离线PV每区至少10点、在线输入至少8点、完整标签12点，三者不混用。训练按目标时间隔离并在训练分区拟合缩放器；事后天气与切分前填充限制了严格在线结论。重复时刻排除或拒绝，完整23/25小时日仍待实现。')
replace(199, '负荷、电价和PV输入维度分别为21/17、26/17和28/27；在线口径修正未改变冻结权重与固定测试指标。')
replace(260, '本章依据生产网络定义说明独立负荷、电价分位和PV v2三个模型。负荷与电价采用相同BiGRU—GRU骨干但分别训练和保存；PV v2以TCN、GRU和四头交叉注意力组合预测。')
replace(261, '权重、缩放器与输入列通过加载契约核验，现有资产不变。后续训练另存新实验，优先Colab T4。旧训练硬件记录和发行时刻协变量可用性仍为复现限制。')

# Added paragraphs are anchored to the original structure, preserving its sections.
def after(anchor, text='', style='Normal'):
    node = OxmlElement('w:p')
    anchor._p.addnext(node)
    p = Paragraph(node, anchor._parent)
    p.style = style
    p.text = text
    return p

archive_p = after(original[278], '实际输入另存backend/cache/forecast_inputs下的JSON.gz文件，以内容SHA-256命名，包含天气记录、提供方、采集时间、质量、各任务past/future特征、资产签名、推理完成后的原生成时间与结果。普通天气端点未提供发行批次时，model_run留空并明确原因。归档与6小时重启缓存独立，不因返回缓存产生新归档；保存失败或损坏明确反馈。离线脚本使用归档和原资产重放全部24步，不访问后期天气、不写数据库、不训练。')
after(original[349], '2026年10月3日的真实输入重放同时核对负荷、电价和光伏全部预测步。三任务最大绝对输出差均为0，说明同输入、同资产可以复现该次预测；这个结果与预测对实际标签的误差是两个不同概念。新增归档只能追溯启用后保存的预测，无法补回此前未保存的天气批次。')

# Replace the outdated UI figure with current real captures and add the second mode.
original[293].clear()
original[293].add_run().add_picture(str(ROOT / 'docs/evidence/frontend/frontend-business-redesign-2026-10-03.jpg'), width=Cm(16.2))
original[292]._p.getparent().remove(original[292]._p)
original[293]._p.addnext(original[292]._p)
intro = after(original[292], '深色监控模式用于持续观察曲线与状态，与正式模式复用相同服务、筛选条件及数据标记，如图5-3所示。')
monitor_pic = after(intro)
monitor_pic.add_run().add_picture(str(ROOT / 'docs/evidence/frontend/frontend-monitor-redesign-2026-10-03.jpg'), width=Cm(16.2))
monitor_cap = after(monitor_pic, '图5-3 深色监控模式的负荷预测页面（2026-10-03）')

# Put every existing figure caption after its inline image.
for cap_idx, pic_idx in [(272,273),(320,321),(323,324),(330,331),(340,341),(342,343)]:
    cap, pic = original[cap_idx], original[pic_idx]
    cap._p.getparent().remove(cap._p)
    pic._p.addnext(cap._p)

# Actual current system diagram: code-native diagram, not an invented product screenshot.
from PIL import Image, ImageDraw, ImageFont
diagram = Image.new('RGB', (1800, 1160), 'white')
draw = ImageDraw.Draw(diagram)
font_path = 'C:/Windows/Fonts/simsun.ttc'
font = ImageFont.truetype(font_path, 32)
small = ImageFont.truetype(font_path, 27)
bold = ImageFont.truetype('C:/Windows/Fonts/simhei.ttf', 35)
def box(x,y,w,h,title,lines):
    draw.rounded_rectangle((x,y,x+w,y+h), radius=8, fill='#F5F6F7', outline='#5C6268', width=2)
    draw.text((x+w/2,y+30), title, font=bold, fill='black', anchor='mm')
    for n,line in enumerate(lines):
        draw.text((x+w/2,y+76+n*38), line, font=small, fill='black', anchor='mm')
def arrow(x1,y1,x2,y2):
    draw.line((x1,y1,x2,y2), fill='#444444', width=3)
    import math
    angle=math.atan2(y2-y1,x2-x1)
    pts=[(x2,y2),(x2-16*math.cos(angle-.45),y2-16*math.sin(angle-.45)),(x2-16*math.cos(angle+.45),y2-16*math.sin(angle+.45))]
    draw.polygon(pts, fill='#444444')
box(560,30,680,145,'React / TypeScript本地页面',['正式工作台 / 深色监控；三任务状态与曲线'])
arrow(900,175,900,240)
box(520,240,760,150,'FastAPI：JWT认证与共享在线编排',['启动预加载；请求合并；独立降级与质量标注'])
box(40,440,390,155,'ISO-NE数据',['八区负荷、RT-LMP、估算PV','完整小时标签独立校验'])
box(1370,440,390,155,'Open-Meteo天气',['历史 / 预报协变量','保留来源、采集时间与质量'])
box(560,455,680,140,'并行准备与特征提供器',['美东小时区间；过去与未来输入窗口'])
arrow(900,390,900,455);arrow(430,520,560,520);arrow(1370,520,1240,520)
box(90,670,460,150,'独立负荷模型',['BiGRU—GRU；168h → 24h'])
box(670,670,460,150,'独立电价模型',['BiGRU—GRU；P10 / P50 / P90'])
box(1250,670,460,150,'PV v2模型',['TCN—GRU—Attention；96h → 24h'])
draw.line((320,635,1480,635), fill='#444444', width=3)
arrow(900,595,900,635)
for x in [320,900,1480]: arrow(x,635,x,670)
draw.line((320,850,1480,850), fill='#444444', width=3)
for x in [320,900,1480]: arrow(x,820,x,850)
box(80,915,475,170,'内存快照 / 本地JSON',['最多6小时；保留原生成时间','有效缓存恢复与页面读取'])
box(660,915,480,170,'实际输入JSON.gz归档',['特征、来源、签名与24步输出','本地离线重放'])
box(1245,915,475,170,'MySQL业务记录',['预测快照与完整标签分别保存','在线误差与历史分析'])
for x in [320,900,1480]: arrow(x,850,x,915)
diagram_path = WORK/'system_architecture_current.png'
diagram.save(diagram_path)
original[273].clear()
original[273].add_run().add_picture(str(diagram_path), width=Cm(16.2))
for pic_idx, name in [(247,'load_price_architecture_current.png'),(324,'load_baseline_current.png'),(341,'pv_baseline_current.png'),(343,'pv_validation_test_current.png')]:
    asset=WORK/name
    if asset.exists():
        original[pic_idx].clear()
        original[pic_idx].add_run().add_picture(str(asset),width=Cm(16.2))

def rewrite_table(table, rows):
    while len(table.rows)>len(rows):
        table._tbl.remove(table.rows[-1]._tr)
    while len(table.rows)<len(rows):
        table.add_row()
    for row, vals in zip(table.rows,rows):
        for cell,text in zip(row.cells,vals): cell.text=text

rewrite_table(tables[7], [
['类别','环境或约定','来源与边界'],
['本地交付','Windows本地；PowerShell启动前后端；MySQL','start.ps1、stop.ps1；不继续云部署'],
['用户开发环境','Windows 11；WSL2 Ubuntu 22.04；Python 3.10；RTX 3050 Laptop 4GB','用户确认TensorFlow识别GPU；本次编辑未重新验证'],
['实际新增诊断','Windows；Python 3.11；TensorFlow 2.16.1；CPU；batch 32','2026-10-03记录；完整离线评估22.922 s，含加载与输出，不含TF导入'],
['现有训练记录','TensorFlow/Keras；已记录GPU:0','精确历史版本、GPU型号及CUDA/cuDNN记录不完整'],
['后续正式训练','Google Colab T4优先；另存.keras、history、metrics','环境安排，不代表旧模型均由T4训练'],
['验证工具','pytest、Vitest、TypeScript、生产构建；真实接口与浏览器','按专项批次报告，不等于本轮重跑全量测试'],
])
rewrite_table(tables[14], [
['日期 / 范围','已有结果','验证边界'],
['10-01 启动与首次加载','后端64项、前端13项通过；真实数据24点；重启有效缓存恢复','5秒探测粒度；不同外网条件，不作严格整机启动对照'],
['10-01 标签与失败恢复','后端68项、前端16项通过；9月30日24个官方电价标签；中断后35.286 s恢复','覆盖跨月、缺标签、日期竞态；不等于长期断网'],
['10-03 输入归档与评估','后端及离线82项、前端5项通过；三任务24步原输入重放差均0','仅说明该输入可复现；未训练或替换权重'],
['10-03 双主题前端','13项主题/导航用例；1440×900、390×844实拍检查通过','主要页面、焦点与菜单；未测帧率和全部浏览器'],
['类型与构建','上述阶段TypeScript检查与生产构建通过','不同专项不相加为全量测试数'],
['认证与噪声','无令牌业务401；合法令牌200；favicon等探测204','仅对精确无关GET/HEAD路径处理'],
['未完成范围','长期连续运行、真实长时断供、完整23/25小时日、估计输入独立精度','明确保留为限制，不列为已通过'],
])
tables[4].rows[5].cells[1].text = '完整标签配对、提前量分组、日期回测、实际输入归档与离线重放'
tables[4].rows[7].cells[1].text = '预加载、有效缓存恢复、共享请求、独立降级、双主题、日志与健康检查'
tables[6].rows[14].cells[3].text = '总体与1—6/7—12/13—24小时负荷误差、输入质量组和排除原因'
for ri in [4,5]: tables[6].rows[ri].cells[2].text='force_refresh（可选）'
api_paths={
14:'/api/analytics/accuracy/stats\n/api/analytics/accuracy/recent\n/api/analytics/drift/check',
15:'/api/analytics/quality/stats\n/api/analytics/quality/validate',
16:'/api/analytics/comparison/models\n/api/analytics/temporal/trends\n/api/analytics/error/distribution\n/api/analytics/patterns/load',
17:'/api/analytics/report/comprehensive\n/api/analytics/dashboard/metrics\n/api/analytics/prediction-vs-actual',
20:'/api/system/status\n/api/system/metrics',
21:'/api/health\n/api/health/liveness\n/api/health/readiness',
22:'/api/auth/register\n/api/auth/login\n/api/auth/refresh\n/api/auth/logout\n/api/auth/logout-all\n/api/auth/change-password',
23:'/api/auth/me\n/api/auth/sessions\n/api/auth/login-history\n/api/auth/session/{session_id}\n/api/auth/health',
}
for ri,text in api_paths.items(): tables[6].rows[ri].cells[1].text=text

# Add compact diagnostic tables without renumbering the established experiment tables.
def insert_table_after(anchor, caption, rows):
    cap=after(anchor,caption)
    t=doc.add_table(rows=1, cols=len(rows[0]))
    rewrite_table(t,rows)
    t._tbl.getparent().remove(t._tbl)
    cap._p.addnext(t._tbl)
    return cap,t

load_cap, load_t = insert_table_after(original[356], '表6-9 负荷测试分层诊断（原模型；重叠目标点）', [
['样本组','目标点数','MAE / MW','条件说明'],
['全部测试','17328','305.40','722窗口；745个不同小时'],
['高温','1182','733.65','温度≥86 °F；事后协变量'],
['节假日','576','538.15','沿用冻结日历字段'],
['非节假日','16752','297.39','与节假日组按相同指标计算'],
['低温','0','—','≤32 °F；无样本，不计算误差'],
])
price_cap, price_t = insert_table_after(original[357], '表6-10 电价区间候选后段验证结果（P50不变）', [
['候选','覆盖率 / %','平均宽度','区间评分','启用结论'],
['原区间','72.08','20.319','35.949','保留生产'],
['全局扩宽','90.58','32.769','38.395','未通过筛选'],
['分提前量扩宽','90.42','32.633','38.390','未通过筛选'],
])
note=after(price_cap,'')
note._p.getparent().remove(note._p)
price_note = Paragraph(OxmlElement('w:p'), original[357]._parent)
price_t._tbl.addnext(price_note._p)
price_note.text = '注：后段验证361个窗口；宽度和区间评分单位为USD/MWh，评分越小越好。校准前段与后段无共享目标小时，独立测试不用于候选启用。'
pv_cap, pv_t = insert_table_after(original[358], '表6-11 光伏完整当地日诊断（25日；600个不同小时）', [
['指标','数值','口径'],
['平均日能量相对误差','15.5244%','完整美东日；ISO-NE估算参考'],
['平均峰值绝对误差','633.8880 MW','每日峰值幅度'],
['平均峰时绝对偏移','0.9200 h','每日峰值小时；不等于全部窗口'],
])

# Rebuild an actual Word TOC field instead of a fixed list of PAGEREF entries.
for i in range(16,86):
    if not original[i]._p.xpath('./w:pPr/w:sectPr'):
        original[i]._p.getparent().remove(original[i]._p)
toc_p = after(original[15])
def field(paragraph, instr, cached=''):
    r=OxmlElement('w:r'); begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');r.append(begin);paragraph._p.append(r)
    r=OxmlElement('w:r'); code=OxmlElement('w:instrText');code.set(qn('xml:space'),'preserve');code.text=instr;r.append(code);paragraph._p.append(r)
    r=OxmlElement('w:r'); sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');r.append(sep);paragraph._p.append(r)
    paragraph.add_run(cached)
    r=OxmlElement('w:r'); end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');r.append(end);paragraph._p.append(r)
field(toc_p, ' TOC \\o "1-3" \\h \\z \\u ', '')
settings=doc.settings.element
up=settings.find(qn('w:updateFields'))
if up is None: up=OxmlElement('w:updateFields');settings.append(up)
up.set(qn('w:val'),'true')

# Remove the artificial blank-page break while preserving section boundaries.
if original[256]._p.getparent() is not None:
    original[256]._p.getparent().remove(original[256]._p)
original[354].paragraph_format.page_break_before=False

# Update only references that were actually checked this turn.
for p in original[382:395]:
    if re.match(r'^\[(8|10|11|12)\]',p.text):
        p.text=p.text.replace('2026-09-13','2026-10-03')

def font_run(run, size=12, chinese='宋体', bold=False):
    run.font.name='Times New Roman';run.font.size=Pt(size);run.font.bold=bold
    run.font.color.rgb=RGBColor(0,0,0)
    rpr=run._r.get_or_add_rPr();rf=rpr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts');rpr.insert(0,rf)
    for key in ['asciiTheme','eastAsiaTheme','hAnsiTheme','cstheme']:
        rf.attrib.pop(qn('w:'+key),None)
    for key,val in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia',chinese),('cs','Times New Roman')]: rf.set(qn('w:'+key),val)

def style_font(style,size=12,chinese='宋体',bold=False):
    style.font.name='Times New Roman';style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
    rpr=style.element.get_or_add_rPr()
    rf=rpr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts');rpr.insert(0,rf)
    for a in list(rf.attrib): rf.attrib.pop(a)
    for key,val in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia',chinese)]: rf.set(qn('w:'+key),val)

normal=doc.styles['Normal'];style_font(normal)
nf=normal.paragraph_format;nf.first_line_indent=Pt(24);nf.line_spacing=1.5;nf.space_before=Pt(0);nf.space_after=Pt(0);nf.widow_control=True
nf.keep_with_next=False;nf.keep_together=False
for name,size in [('Heading 1',16),('Heading 2',14),('Heading 3',12)]:
    st=doc.styles[name];style_font(st,size,'黑体',True)
    pf=st.paragraph_format;pf.first_line_indent=Pt(0);pf.keep_with_next=True;pf.keep_together=True
    pf.line_spacing=1.5;pf.space_before=Pt(12 if name!='Heading 1' else 0);pf.space_after=Pt(6 if name!='Heading 1' else 18)
    pf.page_break_before=(name=='Heading 1');pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if name=='Heading 1' else WD_ALIGN_PARAGRAPH.LEFT
for name,size in [('论文前置标题',16),('论文目录标题',16)]:
    st=doc.styles[name];style_font(st,size,'黑体',True)
    st.paragraph_format.page_break_before=True;st.paragraph_format.first_line_indent=Pt(0);st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;st.paragraph_format.space_after=Pt(18)
style_font(doc.styles['Title'],20,'黑体',True)
for name in ['Figure Caption','Table Caption','Table Text','Table Note','Equation']:
    if name not in doc.styles: doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    st=doc.styles[name];st.base_style=normal;style_font(st,10.5 if name!='Equation' else 12)
    pf=st.paragraph_format;pf.first_line_indent=Pt(0);pf.line_spacing=1.15;pf.space_before=Pt(4);pf.space_after=Pt(6);pf.keep_together=True
    pf.keep_with_next=(name=='Table Caption');pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if 'Caption' in name else WD_ALIGN_PARAGRAPH.LEFT
style_font(doc.styles['Table Caption'],10.5,'宋体',True)
for i in range(1,4):
    name=f'TOC {i}'
    st=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    style_font(st,12,'宋体',i==1)
    pf=st.paragraph_format;pf.first_line_indent=Pt(0);pf.left_indent=Pt((i-1)*12);pf.line_spacing=1.25;pf.space_before=Pt(0);pf.space_after=Pt(3);pf.keep_with_next=False

cover_set={id(p._p) for p in original[:7]}
refs_set={id(p._p) for p in original[382:395]}
for p in doc.paragraphs:
    if p._p.xpath('./w:pPr/w:sectPr'):
        # Section-break paragraphs must remain but take minimal space.
        p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.keep_with_next=False
        continue
    pf=p.paragraph_format
    if id(p._p) in cover_set: continue
    heading=p.style.name.startswith('Heading')
    if p._p.xpath('.//w:drawing'):
        pf.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.first_line_indent=Pt(0)
        pf.space_before=Pt(6);pf.space_after=Pt(0);pf.keep_with_next=True;pf.keep_together=True;pf.line_spacing=1
        # Keep every inline figure within the printable width.
        for inline in p._p.xpath('.//wp:inline'):
            extent=inline.find(qn('wp:extent'))
            if extent is not None and int(extent.get('cx'))>int(Cm(16.2)):
                ratio=int(Cm(16.2))/int(extent.get('cx'))
                extent.set('cx',str(int(Cm(16.2))));extent.set('cy',str(int(int(extent.get('cy'))*ratio)))
                for x in inline.xpath('.//a:xfrm/a:ext'):
                    x.set('cx',extent.get('cx'));x.set('cy',extent.get('cy'))
        continue
    if re.match(r'^图\d+-\d+\s',p.text):
        p.style='Figure Caption';pf.keep_with_next=False;pf.page_break_before=False
        pf.first_line_indent=Pt(0);pf.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.line_spacing=1.15;pf.space_before=Pt(4);pf.space_after=Pt(8)
        for r in p.runs: font_run(r,10.5)
    elif re.match(r'^表\d+-\d+\s',p.text):
        p.style='Table Caption';pf.keep_with_next=True;pf.page_break_before=False
        pf.first_line_indent=Pt(0);pf.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.line_spacing=1.15;pf.space_before=Pt(6);pf.space_after=Pt(4)
        for r in p.runs: font_run(r,10.5,bold=True)
    elif heading:
        size=16 if p.style.name=='Heading 1' else 14 if p.style.name=='Heading 2' else 12
        pf.first_line_indent=Pt(0);pf.keep_with_next=True;pf.keep_together=True
        pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if size==16 else WD_ALIGN_PARAGRAPH.LEFT
        pf.page_break_before=(size==16)
        pf.line_spacing=1.5;pf.space_before=Pt(0 if size==16 else 12);pf.space_after=Pt(18 if size==16 else 6)
        for r in p.runs: font_run(r,size,'黑体',True)
    elif p.style.name in ['论文前置标题','论文目录标题']:
        pf.first_line_indent=Pt(0);pf.keep_with_next=True;pf.line_spacing=1.5
        for r in p.runs: font_run(r,16,'黑体',True)
    elif id(p._p) in refs_set:
        pf.first_line_indent=Pt(-21);pf.left_indent=Pt(21);pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;pf.line_spacing=1.25;pf.space_after=Pt(5);pf.keep_with_next=False;pf.keep_together=True
        for r in p.runs: font_run(r,10.5)
    elif p is price_note or p.text.startswith('注：后段验证'):
        p.style='Table Note';pf.keep_with_next=False;pf.line_spacing=1.15;pf.first_line_indent=Pt(0)
        for r in p.runs:font_run(r,10.5)
    elif p.text:
        pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;pf.first_line_indent=Pt(24);pf.space_before=Pt(0);pf.space_after=Pt(0);pf.line_spacing=1.5;pf.keep_with_next=False;pf.keep_together=False;pf.widow_control=True
        for r in p.runs: font_run(r,12)

# Compact, native editable OMML equations with aligned equation numbers.
equations={219:('zₜ = asinh(pₜ / 50)','4-1'),221:('p̂ₜ = 50 × sinh(ẑₜ)','4-2'),224:('Lτ(y, q) = max[τ(y − q), (τ − 1)(y − q)]，τ ∈ {0.1, 0.5, 0.9}','4-3'),226:('q₁₀ = m − softplus(a)，q₅₀ = m，q₉₀ = m + softplus(b)','4-4'),241:('hδ(e) = {0.5e²（|e| ≤ δ）；δ(|e| − 0.5δ)（|e| > δ）}，δ = 0.05','4-5'),243:('Lpoint = Σ(w × hδ) / max(Σw, 1)','4-6'),245:('L = Lpoint + 0.12Lenergy + 0.08Lpeak','4-7')}
for i,(expr,num) in equations.items():
    p=original[i];p.clear();p.style='Equation'
    pf=p.paragraph_format;pf.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.first_line_indent=Pt(0);pf.keep_with_next=False;pf.keep_together=True;pf.line_spacing=1.15;pf.space_before=Pt(6);pf.space_after=Pt(6)
    math=OxmlElement('m:oMath');mr=OxmlElement('m:r');mt=OxmlElement('m:t');mt.text=expr;mr.append(mt);math.append(mr);p._p.append(math)
    p.add_run('  （'+num+'）')
    for r in p.runs:font_run(r,10.5)

# Keep scientific exponents together in the model tables.
for t in [tables[3],tables[8]]:
    for row in t.rows:
        for c in row.cells:
            c.text=re.sub(r'(\d+(?:\.\d+)?)e-(\d+)',lambda m:m.group(1)+'×10⁻'+m.group(2).translate(str.maketrans('0123456789','⁰¹²³⁴⁵⁶⁷⁸⁹')),c.text)

# Tables: clear compact headers, deliberate widths, repeated headings, no split rows.
width_map={1:[2.9,4.0,4.4,5.0],2:[2.7,2.1,2.1,4.9,4.5],3:[2.9,1.6,4.1,3.2,4.5],4:[2.5,6.5,7.3],5:[3.1,3.5,5.7,4.0],6:[2.0,5.4,3.3,5.6],7:[2.7,7.0,6.6],8:[2.7,5.2,3.2,5.2],9:[4.4,4.2,7.7],10:[5.1,3.8,3.8,3.6],11:[4.4,4.2,7.7],12:[7.0,4.65,4.65],13:[6.0,3.6,3.6,3.1],14:[3.1,7.3,5.9],15:[3.2,2.5,2.6,8.0],16:[4.2,2.7,2.6,2.6,4.2],17:[6.0,3.9,6.4]}
for ti,t in enumerate(doc.tables):
    if ti==0: continue
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=width_map.get(ti,[16.3/len(t.columns)]*len(t.columns))
    # Avoid assumptions if an established table has a different number of columns.
    if len(widths)!=len(t.columns): widths=[16.3/len(t.columns)]*len(t.columns)
    for col,w in zip(t.columns,widths): col.width=Cm(w)
    pr=t._tbl.tblPr
    borders=pr.find(qn('w:tblBorders'))
    if borders is None:borders=OxmlElement('w:tblBorders');pr.append(borders)
    for child in list(borders):borders.remove(child)
    for edge in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        for old in trpr.findall(qn('w:trHeight')):trpr.remove(old)
        if trpr.find(qn('w:cantSplit')) is None:trpr.append(OxmlElement('w:cantSplit'))
        if ri==0 and trpr.find(qn('w:tblHeader')) is None:trpr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr();shd=cp.find(qn('w:shd'))
            if shd is None:shd=OxmlElement('w:shd');cp.append(shd)
            shd.set(qn('w:fill'),'EDEFF1' if ri==0 else 'FFFFFF')
            mar=cp.find(qn('w:tcMar'))
            if mar is None:mar=OxmlElement('w:tcMar');cp.append(mar)
            for side,value in [('top','75'),('bottom','75'),('left','85'),('right','85')]:
                x=mar.find(qn('w:'+side))
                if x is None:x=OxmlElement('w:'+side);mar.append(x)
                x.set(qn('w:w'),value);x.set(qn('w:type'),'dxa')
            for p in c.paragraphs:
                p.style='Table Text';pf=p.paragraph_format;pf.first_line_indent=Pt(0);pf.left_indent=Pt(0);pf.right_indent=Pt(0);pf.line_spacing=1.15;pf.space_before=Pt(0);pf.space_after=Pt(0);pf.keep_with_next=(ti in [15,16,17] and ri<len(t.rows)-1);pf.keep_together=True;pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if ri==0 else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:font_run(r,10.5,bold=(ri==0))

# Cover contents and signature blanks are left as supplied, with normalized title fonts.
for i in [0,1,4]:
    p=original[i];p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
    for r in p.runs:font_run(r,20 if i==4 else 18 if i==1 else 18,'黑体',True)
for p in [original[9],original[13]]:
    p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(6)
    for r in p.runs:font_run(r,12)
original[12].paragraph_format.line_spacing=1.25

doc.settings.odd_and_even_pages_header_footer=False
for si,s in enumerate(doc.sections):
    s.top_margin=Cm(2.5);s.bottom_margin=Cm(2.0);s.left_margin=Cm(2.5);s.right_margin=Cm(2.0)
    s.header_distance=Cm(1.2);s.footer_distance=Cm(1.2)
    s.different_first_page_header_footer=(si==0)
    if si>0:
        s.header.is_linked_to_previous=False
        h=s.header.paragraphs[0];h.text='贵阳人文科技学院本科毕业论文';h.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;h.paragraph_format.first_line_indent=Pt(0)
        for r in h.runs:font_run(r,9)
        s.footer.is_linked_to_previous=False
        p=s.footer.paragraphs[0];p.clear();field(p,' PAGE ','1')
    for p in s.footer.paragraphs:
        p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
        for r in p.runs:font_run(r,10.5)
doc.core_properties.title='基于TensorFlow的智能电网负荷预测系统设计与实现'
doc.core_properties.subject='依据2026年10月3日项目工作区与阶段验收记录修订'
doc.save(OUTPUT)
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==source_sha
(WORK/'edit_manifest.json').write_text(json.dumps({'source':str(SOURCE),'source_sha256':source_sha,'output':str(OUTPUT),'updated_paragraphs':changes,'added_figures':['图5-3'],'new_tables':['表6-9','表6-10','表6-11'],'evidence_cutoff':'2026-10-03','models_changed':False},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'output':str(OUTPUT),'source_unchanged':True,'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'figures':len(doc.inline_shapes),'updated_paragraphs':len(changes)},ensure_ascii=False))
