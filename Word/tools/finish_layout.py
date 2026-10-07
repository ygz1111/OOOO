from pathlib import Path
from copy import deepcopy
from docx import Document
import json

ROOT = Path(__file__).resolve().parents[2]
work = ROOT / 'Word/materials/thesis_update'
output = ROOT / 'Word/本科毕业论文_基于TensorFlow的智能电网负荷预测系统设计与实现_项目更新排版版.docx'
doc = Document(output)
changes = []
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                if paragraph.text not in ('GET/POST', 'GET/DELETE'):
                    continue
                old = paragraph.text
                properties = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
                paragraph.clear()
                run = paragraph.add_run(old.replace('/', '/\n'))
                if properties is not None:
                    run._r.insert(0, properties)
                changes.append({'from': old, 'to': paragraph.text})
assert len(changes) == 2, changes
doc.save(output)
(work/'last_layout_fix.json').write_text(json.dumps(changes, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'method_cells_fixed': len(changes)}, ensure_ascii=False))
