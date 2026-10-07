"""Read-only checks of the DOCX files and retained delivery render snapshot.

This does not render documents, compare archived pixels, or write QA records.
"""

from pathlib import Path, PureWindowsPath
import hashlib
import json
import re
from zipfile import ZipFile

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
WORK = ROOT / 'Word/materials/thesis_update'


def verify_delivery():
    manifest = json.loads((WORK / 'edit_manifest.json').read_text(encoding='utf-8'))
    record = json.loads((WORK / 'final_structural_checks.json').read_text(encoding='utf-8'))
    # The saved manifest records its original absolute Windows paths. Resolve
    # filenames in the current project without rewriting that historical record.
    source = ROOT / 'Word' / PureWindowsPath(manifest['source']).name
    output = ROOT / 'Word' / PureWindowsPath(manifest['output']).name
    latest = WORK / 'delivery_checked'

    source_sha = hashlib.sha256(source.read_bytes()).hexdigest()
    output_sha = hashlib.sha256(output.read_bytes()).hexdigest()
    original = Document(source)
    doc = Document(output)
    cover_preserved = bool(original.tables and doc.tables) and (
        [[cell.text for cell in row.cells] for row in doc.tables[0].rows]
        == [[cell.text for cell in row.cells] for row in original.tables[0].rows]
    )
    with ZipFile(output) as package:
        required_parts = {'[Content_Types].xml', 'word/document.xml', 'word/styles.xml'}
        package_valid = required_parts.issubset(package.namelist())
    structure_valid = package_valid and bool(
        doc.sections and doc.tables and any(p.text.strip() for p in doc.paragraphs)
    )

    page_numbers = sorted(
        int(match.group(1))
        for path in latest.glob('page-*.png')
        if (match := re.fullmatch(r'page-(\d+)\.png', path.name))
        and path.is_file() and path.stat().st_size > 0
    )
    expected_pages = int(record['pages'])
    checks = {
        'original_sha256_matches': source_sha == manifest['source_sha256'],
        'cover_fields_preserved': cover_preserved,
        'output_basic_structure_valid': structure_valid,
        'output_matches_recorded_sha256': output_sha == record.get('output_sha256'),
        'retained_render_pages_complete': latest.is_dir() and page_numbers == list(range(1, expected_pages + 1)),
    }
    return {
        'status': 'pass' if all(checks.values()) else 'fail',
        'checks': checks,
        'source': str(source),
        'output': str(output),
        'output_sha256': output_sha,
        'output_bytes': output.stat().st_size,
        'output_structure': {
            'paragraphs': len(doc.paragraphs),
            'tables': len(doc.tables),
            'sections': len(doc.sections),
            'inline_shapes': len(doc.inline_shapes),
        },
        'retained_render_directory': str(latest),
        'retained_render_page_count': len(page_numbers),
        'retained_render_page_numbers': page_numbers,
        'expected_render_page_count_from_record': expected_pages,
        'document_page_count_verified': False,
        'pixel_comparison_performed': False,
        'visual_review_performed': False,
        'limitations': [
            'Only the retained delivery_checked snapshot is available; the old final_checked images were archived.',
            'The saved page count and output SHA are provenance checks, not a new pagination or visual review.',
        ],
    }


if __name__ == '__main__':
    result = verify_delivery()
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result['status'] == 'pass' else 1)
